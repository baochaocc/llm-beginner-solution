"""把实验报告 Markdown 转成 Word 文档（python-docx，无模板依赖）。

用法（在仓库根目录）：
    python scripts/md_to_docx.py 实验报告.md 实验报告.docx

支持语法：`#/##/###` 标题、Markdown 表格、``` 代码块、`- ` 列表、`> ` 引用、
`![图注](相对路径)` 图片（自动插入并居中，附图注）。
字体：正文宋体五号（10.5pt），标题宋体加粗，代码 Consolas。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

IMG_RE = re.compile(r"!\[(.*?)\]\((.*?)\)")
MAX_IMG_WIDTH_CM = 15.5


def set_font(style, name_zh="宋体", size=10.5, bold=None, color=None):
    style.font.name = "Consolas" if name_zh == "Consolas" else "SimSun"
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor(*color)
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), name_zh)


def clean_inline(text):
    text = IMG_RE.sub(r"\1", text)          # 表格/正文里的图片语法退化为图注文字
    return text.replace("**", "").replace("`", "").strip()


def is_table_sep(cells):
    return all(c.strip() and set(c.strip()) <= set(":-") for c in cells)


def add_code_line(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)


def add_image(doc, md_dir: Path, alt: str, src: str):
    img = Path(src)
    if not img.is_absolute():
        img = md_dir / src
    if not img.exists():
        doc.add_paragraph(f"[图片缺失：{src}]")
        return
    try:
        doc.add_picture(str(img), width=Cm(MAX_IMG_WIDTH_CM))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        run = cap.add_run(alt or img.name)
        run.font.size = Pt(9)
        run.font.italic = True
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as e:                             # noqa: BLE001
        doc.add_paragraph(f"[图片插入失败：{src}（{e}）]")


def build(md_path: Path, docx_path: Path):
    doc = Document()
    set_font(doc.styles["Normal"], "宋体", 10.5)
    for style_name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12.5)):
        st = doc.styles[style_name]
        set_font(st, "宋体", size, bold=True, color=(0x1F, 0x3B, 0x63))
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)

    md_dir = md_path.parent
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, n, in_code = 0, len(lines), False

    while i < n:
        raw = lines[i]
        line = raw.strip()

        if line.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            add_code_line(doc, raw)
            i += 1
            continue
        if not line:
            i += 1
            continue

        # 表格
        if line.startswith("|") and i + 1 < n and lines[i + 1].strip().startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not is_table_sep(cells):
                    rows.append([clean_inline(c) for c in cells])
                i += 1
            if rows:
                table = doc.add_table(rows=0, cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_i, row in enumerate(rows):
                    cells = table.add_row().cells
                    for c_i, text in enumerate(row):
                        if c_i < len(cells):
                            cells[c_i].text = ""
                            run = cells[c_i].paragraphs[0].add_run(text)
                            run.font.name = "SimSun"
                            run.font.size = Pt(10)
                            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                            if r_i == 0:
                                run.font.bold = True
            continue

        # 图片
        m = IMG_RE.fullmatch(line)
        if m:
            add_image(doc, md_dir, m.group(1), m.group(2))
            i += 1
            continue

        if line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=3)
        elif line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=2)
        elif line.startswith("# "):
            doc.add_heading(clean_inline(line[2:]), level=1)
        elif line.startswith("- "):
            doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
        elif line.startswith("> "):
            p = doc.add_paragraph()
            run = p.add_run(clean_inline(line[2:]))
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        else:
            doc.add_paragraph(clean_inline(line))
        i += 1

    doc.save(str(docx_path))
    print(f"已生成：{docx_path}（{docx_path.stat().st_size / 1024:.1f} KB）")


if __name__ == "__main__":
    md = Path(sys.argv[1] if len(sys.argv) > 1 else "实验报告.md")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else str(md.with_suffix(".docx")))
    build(md.resolve(), out.resolve())
